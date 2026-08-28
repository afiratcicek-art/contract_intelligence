"""Deterministic Dispute Ready pack — DOCX from Document(), never user-docx.

No LLM, no embeddings, no external model calls. Künye + impacts + chronology + issues
+ exhibit index. Zip-of-PDFs is deferred (TB-59); exhibit rows stay on
the dossier so files remain with the case.
"""
from __future__ import annotations

import logging
from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Optional

from docx import Document
from docx.shared import Pt

from backend.utils.file_handler import upload_document

logger = logging.getLogger(__name__)

_ENTITY_TABLE = {
    "change": ("changes", "change_number", "title"),
    "correspondence": ("correspondences", "corr_number", "subject"),
    "rfi": ("rfis", "rfi_number", "subject"),
}


def _heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _para(doc: Document, text: str, *, italic: bool = False) -> None:
    p = doc.add_paragraph(text or "—")
    if italic and p.runs:
        p.runs[0].italic = True
    for run in p.runs:
        run.font.size = Pt(11)


def _kv(doc: Document, label: str, value: Any) -> None:
    text = "—" if value is None or value == "" else str(value)
    p = doc.add_paragraph()
    run_l = p.add_run(f"{label}: ")
    run_l.bold = True
    run_l.font.size = Pt(11)
    run_v = p.add_run(text)
    run_v.font.size = Pt(11)


def _lookup_entity(db, ref_type: str, entity_id: Optional[str]) -> str:
    if not entity_id or ref_type not in _ENTITY_TABLE:
        return ""
    table, num_col, title_col = _ENTITY_TABLE[ref_type]
    try:
        res = (
            db.table(table)
            .select(f"{num_col}, {title_col}")
            .eq("id", entity_id)
            .limit(1)
            .execute()
        )
    except Exception as exc:
        logger.warning("exhibit entity lookup failed: %s", exc)
        return ""
    if not res.data:
        return ""
    row = res.data[0]
    num = row.get(num_col) or ""
    title = row.get(title_col) or ""
    return f"{num} — {title}".strip(" —")


def _lookup_document(db, document_id: Optional[str]) -> str:
    if not document_id:
        return ""
    try:
        res = (
            db.table("pdf_document")
            .select("original_filename, storage_path")
            .eq("id", document_id)
            .limit(1)
            .execute()
        )
    except Exception as exc:
        logger.warning("exhibit document lookup failed: %s", exc)
        return ""
    if not res.data:
        return ""
    return res.data[0].get("original_filename") or res.data[0].get("storage_path") or ""


def collect_exhibits(db, dossier: dict) -> list[dict]:
    """Related files stored WITH the report — index only in v1 (no zip)."""
    exhibits: list[dict] = []
    seen: set[str] = set()

    def add(item: dict) -> None:
        key = item.get("key") or f"{item.get('kind')}:{item.get('label')}"
        if key in seen:
            return
        seen.add(key)
        exhibits.append(item)

    dispute_id = dossier.get("id")
    try:
        docs = (
            db.table("pdf_document")
            .select("id, original_filename, storage_path")
            .eq("entity_type", "dispute")
            .eq("entity_id", dispute_id)
            .execute()
        )
        for row in docs.data or []:
            add({
                "key": f"pdf:{row.get('id')}",
                "kind": "document",
                "label": row.get("original_filename") or row.get("storage_path") or row.get("id"),
                "storage_path": row.get("storage_path"),
            })
    except Exception as exc:
        logger.warning("dispute pdf_document list failed: %s", exc)

    for issue in dossier.get("dispute_issues") or []:
        for pos in issue.get("dispute_positions") or []:
            for ref in pos.get("dispute_position_refs") or []:
                ref_type = ref.get("ref_type") or ""
                if ref_type == "manual":
                    add({
                        "key": f"manual:{ref.get('id')}",
                        "kind": "manual",
                        "label": ref.get("manual_title") or "manual",
                        "note": ref.get("manual_note"),
                        "date": ref.get("manual_date"),
                    })
                    continue
                label = _lookup_entity(db, ref_type, ref.get("entity_id"))
                doc_label = _lookup_document(db, ref.get("document_id"))
                add({
                    "key": f"ref:{ref.get('id')}",
                    "kind": ref_type,
                    "label": label or doc_label or (ref.get("entity_id") or ref.get("document_id") or ""),
                    "document_id": ref.get("document_id"),
                    "entity_id": ref.get("entity_id"),
                })
    return exhibits


def build_pack_docx(dossier: dict, chronology: Optional[dict], exhibits: list[dict]) -> bytes:
    doc = Document()
    number = dossier.get("dispute_number") or ""
    title = dossier.get("title") or ""
    _heading(doc, f"{number} — {title}", 0)

    _heading(doc, "1. Künye", 1)
    _kv(doc, "Number", number)
    _kv(doc, "Title", title)
    _kv(doc, "Status", dossier.get("status"))
    _kv(doc, "Origin", dossier.get("origin"))
    _kv(doc, "Venue", dossier.get("venue"))
    _kv(doc, "Summary", dossier.get("summary"))
    _kv(doc, "Source change", dossier.get("source_change_id"))
    _kv(doc, "Source correspondence", dossier.get("source_correspondence_id"))
    _kv(doc, "Created", dossier.get("created_at"))

    _heading(doc, "2. Impacts", 1)
    impacts = sorted(
        dossier.get("dispute_impacts") or [],
        key=lambda r: (r.get("sort_order") or 0),
    )
    if not impacts:
        _para(doc, "No impacts recorded.", italic=True)
    for row in impacts:
        amount = row.get("amount")
        unit = row.get("unit") or ""
        amt = f"{amount} {unit}".strip() if amount is not None else "—"
        _para(doc, f"[{row.get('type')}] {row.get('label')}: {amt}")
        if row.get("notes"):
            _para(doc, str(row["notes"]), italic=True)

    _heading(doc, "3. Chronology", 1)
    events = []
    if chronology:
        events = [
            e for e in (chronology.get("events") or [])
            if e.get("is_active", True)
        ]
        events.sort(key=lambda e: e.get("event_date") or "")
    if not events:
        _para(doc, "No chronology events.", italic=True)
    for ev in events:
        subj = ev.get("subject") or ev.get("approved_narrative") or ev.get("auto_narrative") or ""
        _para(
            doc,
            f"{ev.get('event_date') or '—'}  [{ev.get('event_type')}]  {subj}".rstrip(),
        )

    _heading(doc, "4. Issues / claims / responses", 1)
    issues = sorted(
        dossier.get("dispute_issues") or [],
        key=lambda r: (r.get("sort_order") or 0),
    )
    if not issues:
        _para(doc, "No disputed issues recorded.", italic=True)
    for issue in issues:
        _heading(doc, str(issue.get("title") or "Issue"), 2)
        positions = sorted(
            issue.get("dispute_positions") or [],
            key=lambda r: (0 if r.get("side") == "claim" else 1, r.get("sort_order") or 0),
        )
        for pos in positions:
            _para(doc, f"{(pos.get('side') or '').upper()}: {pos.get('title')}")
            if pos.get("summary"):
                _para(doc, str(pos["summary"]), italic=True)
            for ref in pos.get("dispute_position_refs") or []:
                if ref.get("ref_type") == "manual":
                    _para(doc, f"  — manual: {ref.get('manual_title')}")
                else:
                    _para(
                        doc,
                        f"  — {ref.get('ref_type')}: "
                        f"{ref.get('entity_id') or ref.get('document_id') or ''}",
                    )

    _heading(doc, "5. Exhibit index", 1)
    _para(
        doc,
        "Related files stay with this dossier. v1 lists them here; "
        "a zip-of-PDFs bundle is deferred.",
        italic=True,
    )
    if not exhibits:
        _para(doc, "No exhibits attached.", italic=True)
    for i, item in enumerate(exhibits, start=1):
        extra = item.get("note") or item.get("storage_path") or ""
        line = f"E{i:02d}  [{item.get('kind')}]  {item.get('label')}"
        if extra and extra != item.get("label"):
            line = f"{line}  ({extra})"
        _para(doc, line)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_and_store_pack(
    db,
    dossier: dict,
    chronology: Optional[dict],
    project_id: str,
) -> dict:
    """Build DOCX, upload Storage, return pack_storage_path / pack_generated_at."""
    exhibits = collect_exhibits(db, dossier)
    payload = build_pack_docx(dossier, chronology, exhibits)
    number = dossier.get("dispute_number") or "DSP"
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    filename = f"{number}-pack-{stamp}.docx"
    path = upload_document(
        payload,
        filename,
        project_id,
        "dispute",
        str(dossier["id"]),
    )
    generated_at = datetime.now(timezone.utc).isoformat()
    return {
        "pack_storage_path": path,
        "pack_generated_at": generated_at,
        "exhibit_count": len(exhibits),
    }
