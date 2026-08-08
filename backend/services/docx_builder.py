"""Clean-slate DOCX builder for authored drafts.

Never ingests a user-supplied .docx — always starts from Document().
Chrome layout (field_config.chrome) is export-faithful to Config letterhead:
width_pct, align/text_align, stack, band_height→margins, offset nudge,
watermark as behind-text page image with baked opacity.
"""
from __future__ import annotations

import logging
import re
from html.parser import HTMLParser
from io import BytesIO
from typing import Any, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Mm, Pt

from backend.utils.file_handler import download_document

logger = logging.getLogger(__name__)

# A4 usable content width ≈ 210mm − 25mm×2 margins → ~160mm ≈ 6.3"
_USABLE_WIDTH_IN = 6.3
_A4_WIDTH_MM = 210
_A4_HEIGHT_MM = 297


def _chrome_slot(tpl: dict, slot: str) -> dict:
    cfg = tpl.get("field_config") if isinstance(tpl.get("field_config"), dict) else {}
    chrome = cfg.get("chrome") if isinstance(cfg.get("chrome"), dict) else {}
    slot_cfg = chrome.get(slot) if isinstance(chrome.get(slot), dict) else {}
    return slot_cfg


def _clamp(n: float, lo: float, hi: float) -> float:
    return max(lo, min(hi, n))


def _para_align(slot_cfg: dict, *, text: bool = False):
    key = "text_align" if text else "align"
    align = str(slot_cfg.get(key) or slot_cfg.get("align") or "center").lower()
    return {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }.get(align, WD_ALIGN_PARAGRAPH.CENTER)


def _width_pct(slot_cfg: dict, default: float = 42.0) -> float:
    try:
        return _clamp(float(slot_cfg.get("width_pct", default)), 8.0, 100.0)
    except (TypeError, ValueError):
        return default


def _band_height_inches(slot_cfg: dict, default_px: float = 140.0) -> float:
    """Config band_height_px (≈96dpi UI) → inches."""
    try:
        px = float(slot_cfg.get("band_height_px", default_px))
    except (TypeError, ValueError):
        px = default_px
    return _clamp(px, 72.0, 320.0) / 96.0


def _picture_size_inches(
    img_bytes: bytes,
    slot_cfg: dict,
    *,
    max_height_in: Optional[float] = None,
) -> tuple[float, float]:
    """Width from width_pct; height preserves aspect; both capped to band."""
    from PIL import Image

    target_w = _USABLE_WIDTH_IN * (_width_pct(slot_cfg) / 100.0)
    try:
        im = Image.open(BytesIO(img_bytes))
        im.load()
        iw, ih = im.size
        if iw <= 0 or ih <= 0:
            return target_w, target_w * 0.4
        aspect = ih / float(iw)
    except Exception:
        return target_w, target_w * 0.4

    w, h = target_w, target_w * aspect
    cap = (
        max_height_in
        if max_height_in is not None
        else _band_height_inches(slot_cfg) * 0.72
    )
    if h > cap > 0:
        scale = cap / h
        w, h = w * scale, cap
    return w, h


def _apply_offset_indent(paragraph, slot_cfg: dict) -> None:
    """Map offset_x_pct (−45…45) → paragraph indent nudge."""
    try:
        ox = float(slot_cfg.get("offset_x_pct") or 0)
    except (TypeError, ValueError):
        ox = 0.0
    ox = _clamp(ox, -45.0, 45.0)
    if abs(ox) < 0.5:
        return
    nudge_in = (_USABLE_WIDTH_IN * ox) / 100.0
    pf = paragraph.paragraph_format
    if ox > 0:
        pf.left_indent = Inches(nudge_in)
    else:
        pf.right_indent = Inches(abs(nudge_in))


def _next_paragraph(container):
    """Prefer empty first header/footer paragraph; else append."""
    if container.paragraphs:
        p0 = container.paragraphs[0]
        if not p0.runs and not (p0.text or "").strip():
            return p0
    return container.add_paragraph()


def _fill_band(
    container,
    *,
    img_bytes: Optional[bytes],
    text: Optional[str],
    slot_cfg: dict,
) -> None:
    """Header/footer: stack, align, size, offset — matches Config chrome."""
    stack = str(slot_cfg.get("stack") or "image_first")
    max_h = _band_height_inches(slot_cfg) * 0.75
    first = True

    def add_image() -> None:
        nonlocal first
        if not img_bytes:
            return
        try:
            w, h = _picture_size_inches(img_bytes, slot_cfg, max_height_in=max_h)
            p = _next_paragraph(container) if first else container.add_paragraph()
            first = False
            run = p.add_run()
            run.add_picture(BytesIO(img_bytes), width=Inches(w), height=Inches(h))
            p.alignment = _para_align(slot_cfg)
            _apply_offset_indent(p, slot_cfg)
        except Exception as exc:
            logger.warning("Band image embed failed: %s", exc)

    def add_text() -> None:
        nonlocal first
        if not text:
            return
        p = _next_paragraph(container) if first else container.add_paragraph()
        first = False
        p.add_run(str(text))
        p.alignment = _para_align(slot_cfg, text=True)

    if stack == "text_first":
        add_text()
        add_image()
    else:
        add_image()
        add_text()


def _watermark_bytes_with_opacity(img_bytes: bytes, opacity: float) -> bytes:
    """Bake Config opacity into PNG alpha for Word."""
    from PIL import Image

    op = _clamp(float(opacity), 0.05, 0.45)
    im = Image.open(BytesIO(img_bytes))
    im.load()
    if im.mode != "RGBA":
        im = im.convert("RGBA")
    r, g, b, a = im.split()
    a = a.point(lambda x: int(x * op))
    out = Image.merge("RGBA", (r, g, b, a))
    buf = BytesIO()
    out.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


def _add_page_watermark(section, img_bytes: bytes, slot_cfg: dict) -> None:
    """Floating image behind text, page-anchored (Word watermark)."""
    from docx.oxml import parse_xml

    try:
        opacity = float(slot_cfg.get("opacity") or 0.14)
    except (TypeError, ValueError):
        opacity = 0.14
    try:
        faded = _watermark_bytes_with_opacity(img_bytes, opacity)
    except Exception as exc:
        logger.warning("Watermark opacity bake failed, using original: %s", exc)
        faded = img_bytes

    w_in, h_in = _picture_size_inches(faded, slot_cfg, max_height_in=4.5)

    try:
        oy = float(slot_cfg.get("offset_y_pct") or 0)
    except (TypeError, ValueError):
        oy = 0.0
    oy = _clamp(oy, -40.0, 40.0)

    header = section.header
    p = header.add_paragraph()
    run = p.add_run()
    inline_shape = run.add_picture(
        BytesIO(faded), width=Inches(w_in), height=Inches(h_in)
    )

    try:
        ct_inline = inline_shape._inline
        graphics = ct_inline.xpath("./a:graphic")
        if not graphics:
            return
        graphic = graphics[0]
        cx = int(ct_inline.extent.cx)
        cy = int(ct_inline.extent.cy)
        align_h = str(slot_cfg.get("align") or "center").lower()
        if align_h not in ("left", "right", "center"):
            align_h = "center"

        page_h_emu = int(Mm(_A4_HEIGHT_MM))
        if abs(oy) < 1:
            pos_v = (
                '<wp:positionV relativeFrom="page">'
                "<wp:align>center</wp:align>"
                "</wp:positionV>"
            )
        else:
            center_y = int(page_h_emu / 2 - cy / 2 + (oy / 100.0) * page_h_emu * 0.5)
            center_y = max(0, center_y)
            pos_v = (
                f'<wp:positionV relativeFrom="page">'
                f"<wp:posOffset>{center_y}</wp:posOffset>"
                f"</wp:positionV>"
            )

        anchor_xml = (
            '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/'
            'wordprocessingDrawing" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'simplePos="0" relativeHeight="0" behindDoc="1" locked="0" '
            'layoutInCell="1" allowOverlap="1">'
            '<wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="page"><wp:align>{align_h}</wp:align></wp:positionH>'
            f"{pos_v}"
            f'<wp:extent cx="{cx}" cy="{cy}"/>'
            '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
            "<wp:wrapNone/>"
            '<wp:docPr id="9001" name="Watermark" descr="letterhead watermark"/>'
            "<wp:cNvGraphicFramePr/>"
            "</wp:anchor>"
        )
        ct_anchor = parse_xml(anchor_xml)
        ct_anchor.append(graphic)
        ct_drawing = ct_inline.getparent()
        for child in list(ct_drawing):
            if child.tag.endswith("inline"):
                ct_drawing.remove(child)
        ct_drawing.append(ct_anchor)
    except Exception as exc:
        logger.warning(
            "Watermark anchor conversion failed — inline header image kept: %s",
            exc,
        )


def _configure_a4_section(section, header_cfg: dict, footer_cfg: dict) -> None:
    section.page_width = Mm(_A4_WIDTH_MM)
    section.page_height = Mm(_A4_HEIGHT_MM)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)

    h_band = _band_height_inches(header_cfg, 150)
    f_band = _band_height_inches(footer_cfg, 120)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)
    section.top_margin = Inches(_clamp(0.55 + h_band * 0.45, 0.75, 2.0))
    section.bottom_margin = Inches(_clamp(0.55 + f_band * 0.45, 0.75, 1.8))


def _download_storage_bytes(
    storage_path: Optional[str], project_id: Optional[str]
) -> Optional[bytes]:
    if not storage_path or not project_id:
        return None
    try:
        return download_document(storage_path, project_id)
    except Exception as exc:
        logger.warning("Chrome image fetch failed: %s | path=%s", exc, storage_path)
        return None


class _HtmlToDocx(HTMLParser):
    """Minimal HTML → python-docx walker for B2-supported tags."""

    def __init__(self, doc: Document):
        super().__init__()
        self.doc = doc
        self._p = None
        self._bold = False
        self._italic = False
        self._in_list = None  # 'ul' | 'ol' | None
        self._list_index = 0
        self._skip = False

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag in ("script", "style", "iframe"):
            self._skip = True
            return
        if tag in ("p", "h1", "h2", "h3", "h4"):
            self._p = self.doc.add_paragraph()
            if tag.startswith("h"):
                run = self._p.add_run("")
                size = {"h1": 16, "h2": 14, "h3": 12, "h4": 11}.get(tag, 11)
                run.font.size = Pt(size)
                run.bold = True
        elif tag == "br":
            if self._p is None:
                self._p = self.doc.add_paragraph()
            self._p.add_run().add_break()
        elif tag == "strong" or tag == "b":
            self._bold = True
        elif tag in ("em", "i"):
            self._italic = True
        elif tag in ("ul", "ol"):
            self._in_list = tag
            self._list_index = 0
        elif tag == "li":
            self._list_index += 1
            self._p = self.doc.add_paragraph(
                style="List Bullet" if self._in_list == "ul" else "List Number"
            )
        elif tag == "table":
            self._table_rows = []
            self._current_row = None
            self._in_table = True
        elif tag == "tr":
            self._current_row = []
        elif tag in ("td", "th"):
            self._cell_buf = []
            self._in_cell = True

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in ("script", "style", "iframe"):
            self._skip = False
            return
        if tag in ("strong", "b"):
            self._bold = False
        elif tag in ("em", "i"):
            self._italic = False
        elif tag in ("ul", "ol"):
            self._in_list = None
        elif tag in ("p", "h1", "h2", "h3", "h4", "li"):
            self._p = None
        elif tag in ("td", "th"):
            text = "".join(self._cell_buf)
            if self._current_row is not None:
                self._current_row.append(text)
            self._in_cell = False
        elif tag == "tr":
            if self._current_row is not None:
                self._table_rows.append(self._current_row)
            self._current_row = None
        elif tag == "table":
            rows = getattr(self, "_table_rows", [])
            if rows:
                cols = max(len(r) for r in rows)
                table = self.doc.add_table(rows=len(rows), cols=cols)
                for i, row in enumerate(rows):
                    for j, cell_text in enumerate(row):
                        table.rows[i].cells[j].text = cell_text
            self._in_table = False

    def handle_data(self, data):
        if self._skip:
            return
        if getattr(self, "_in_cell", False):
            self._cell_buf.append(data)
            return
        if not data:
            return
        if self._p is None:
            self._p = self.doc.add_paragraph()
        run = self._p.add_run(data)
        run.bold = self._bold
        run.italic = self._italic


def build_docx(
    *,
    body_html: str,
    field_values: dict[str, Any],
    template: Optional[dict] = None,
) -> bytes:
    """Build a clean A4 .docx from template chrome + fields + body_html."""
    doc = Document()
    section = doc.sections[0]
    tpl = template or {}
    project_id = tpl.get("project_id")

    header_cfg = _chrome_slot(tpl, "header")
    footer_cfg = _chrome_slot(tpl, "footer")
    wm_cfg = _chrome_slot(tpl, "watermark")
    _configure_a4_section(section, header_cfg, footer_cfg)

    wm_bytes = _download_storage_bytes(tpl.get("watermark_image_path"), project_id)
    if wm_bytes:
        try:
            _add_page_watermark(section, wm_bytes, wm_cfg)
        except Exception as exc:
            logger.warning("Watermark embed failed: %s", exc)

    _fill_band(
        section.header,
        img_bytes=_download_storage_bytes(tpl.get("header_image_path"), project_id),
        text=tpl.get("header_text"),
        slot_cfg=header_cfg,
    )
    _fill_band(
        section.footer,
        img_bytes=_download_storage_bytes(tpl.get("footer_image_path"), project_id),
        text=tpl.get("footer_text"),
        slot_cfg=footer_cfg,
    )

    attention = field_values.get("attention_to") or ""
    project = field_values.get("project") or ""

    if attention:
        p = doc.add_paragraph()
        p.add_run("Attention to: ").bold = True
        p.add_run(str(attention))
    if project:
        p = doc.add_paragraph()
        p.add_run("Project: ").bold = True
        p.add_run(str(project))

    doc.add_paragraph()

    parser = _HtmlToDocx(doc)
    safe_html = re.sub(r"\son\w+\s*=", " ", body_html or "", flags=re.I)
    safe_html = re.sub(r"javascript\s*:", "", safe_html, flags=re.I)
    try:
        parser.feed(safe_html)
        parser.close()
    except Exception as exc:
        logger.warning("HTML→docx parse fallback to plain text: %s", exc)
        plain = re.sub(r"<[^>]+>", "", safe_html)
        doc.add_paragraph(plain)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
